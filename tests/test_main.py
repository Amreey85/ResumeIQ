from io import BytesIO

import docx
import pytest

from main import calculate_match_score, extract_text, find_matches, parse_keywords


def test_parse_keywords_normalizes_and_dedupes():
    assert parse_keywords(" Python, machine   learning,,python ") == ["python", "machine learning"]


def test_score_empty_keywords():
    assert calculate_match_score("anything", "") == 0.0


def test_score_partial_match():
    assert calculate_match_score("Senior Python developer", "python, java, go") == 33.3


def test_symbol_keywords_match():
    text = "Skilled in C++, C# and .NET"
    assert find_matches(text, ["c++", "c#", ".net"]) == ["c++", "c#", ".net"]


def test_no_partial_word_match():
    assert find_matches("javascript expert", ["java"]) == []


def test_extract_docx_includes_tables():
    doc = docx.Document()
    doc.add_paragraph("Jane Doe")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Kubernetes"
    buf = BytesIO()
    doc.save(buf)
    text = extract_text("cv.DOCX", buf.getvalue())
    assert "Jane Doe" in text and "Kubernetes" in text


def test_extract_unsupported_type():
    with pytest.raises(ValueError):
        extract_text("cv.txt", b"hello")
